"""
Insert formula images into the IEEE_Formal_Documentation_v2.docx
at the correct locations (after the "Formula" bold text).
"""
import os, sys
sys.stdout.reconfigure(encoding='utf-8')

from docx import Document
from docx.shared import Inches

DOC_PATH = r"c:\Users\ASUS\Documents\Internship\3M\c\IEEE_Formal_Documentation_v2.docx"
OUT_PATH = r"c:\Users\ASUS\Documents\Internship\3M\c\IEEE_Formal_Documentation_v3.docx"
BASE = r"c:\Users\ASUS\Documents\Internship\3M\c"

INSERTIONS = [
    ("Formula", "Joint Angle", "formula_joint_angle.png"),
    ("Formula", "Similarity Score", "formula_similarity_score.png"),
    ("Formula", "Ghost Skeleton Projection", "formula_ghost_skeleton.png"),
]

doc = Document(DOC_PATH)

for _, label, img_file in INSERTIONS:
    img_path = os.path.join(BASE, img_file)
    if not os.path.exists(img_path):
        print(f"WARNING: {img_path} not found, skipping")
        continue
    
    inserted = False
    for i, para in enumerate(doc.paragraphs):
        # Look for the broken image alt text line
        if label in para.text and ("Formula" in para.text or "formula" in para.text.lower()):
            # Check if next paragraph has the broken alt-text from markdown
            if i + 1 < len(doc.paragraphs):
                next_p = doc.paragraphs[i + 1]
                # The converter turns ![alt](url) into just the alt text when image fails
                if any(kw in next_p.text for kw in ["Joint Angle Formula", "Similarity Score Formula", "Ghost Skeleton Alignment Formula"]):
                    next_p.clear()
                    run = next_p.add_run()
                    run.add_picture(img_path, width=Inches(5.0))
                    print(f"OK: Inserted {img_file} replacing alt-text after '{label}'")
                    inserted = True
                    break
            
            # Otherwise just add to this paragraph
            run = para.add_run()
            run.add_break()
            run.add_picture(img_path, width=Inches(5.0))
            print(f"OK: Inserted {img_file} into '{label}' paragraph")
            inserted = True
            break
    
    if not inserted:
        # Brute force: search for any paragraph referencing the image alt text
        for i, para in enumerate(doc.paragraphs):
            if label.split()[0] in para.text and "Formula" in para.text:
                para.clear()
                run = para.add_run()
                run.add_picture(img_path, width=Inches(5.0))
                print(f"OK: Replaced paragraph with {img_file}")
                inserted = True
                break
        if not inserted:
            print(f"SKIP: Could not find insertion point for {img_file}")

doc.save(OUT_PATH)
print(f"\nDone! Saved: {OUT_PATH}")
